"""P2 pure logic: parsing, claim extraction (fake LLM), validation, scoring."""

import io
import os
from datetime import date

os.environ.setdefault("SECRET_KEY", "x" * 32)
os.environ.setdefault("ADMIN_PASSWORD", "admin-pass")

import docx
import pymupdf
import pytest

from app.verification.extract import ClaimIn, extract_claims
from app.verification.parse import ParseError, extract_text
from app.verification.scoring import NO_INFO_MESSAGE, score_review
from app.verification.types import Passage
from app.verification.validate import validate

# ── parsing ───────────────────────────────────────────────────────────────


def test_txt_and_docx_are_parsed() -> None:
    assert extract_text("o.txt", "رأي قانوني في مسألة الإجازات".encode()) == "رأي قانوني في مسألة الإجازات"
    d = docx.Document()
    d.add_paragraph("الفقرة الأولى")
    d.add_paragraph("الفقرة الثانية")
    buf = io.BytesIO()
    d.save(buf)
    assert extract_text("o.docx", buf.getvalue()) == "الفقرة الأولى\nالفقرة الثانية"


def test_pdf_text_layer_is_parsed_and_scans_are_rejected() -> None:
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "Legal opinion about article 41 of the labour law.")
    assert "article 41" in extract_text("o.pdf", doc.tobytes())
    empty = pymupdf.open()
    empty.new_page()
    with pytest.raises(ParseError) as e:
        extract_text("scan.pdf", empty.tobytes())
    assert "ممسوح" in e.value.message_ar


def test_unsupported_type_is_rejected() -> None:
    with pytest.raises(ParseError):
        extract_text("x.exe", b"MZ")


# ── claim extraction ──────────────────────────────────────────────────────


class FakeLLM:
    def __init__(self, *outputs: dict) -> None:
        self.outputs = list(outputs)
        self.calls: list[dict] = []

    async def call_structured(self, system, user, tool_name, schema, *, prompt_version, max_tokens=8192):
        self.calls.append({"system": system, "user": user, "tool": tool_name, "schema": schema})
        return self.outputs.pop(0)


def _claim(text: str, type_: str = "legal_conclusion", mat: str = "core", refs=None) -> dict:
    return {"text_ar": text, "type": type_, "materiality": mat, "cited_refs": refs or []}


@pytest.mark.asyncio
async def test_extract_claims_assigns_ids_dedupes_and_caps() -> None:
    llm = FakeLLM({"claims": [
        _claim("يستحق العامل إجازة ثلاثين يوماً", refs=[{"law_number": "6", "year": 2010, "article": 70, "raw": "م70"}]),
        _claim("يستحق العاملُ إجازة ثلاثين يوماً"),  # same after normalization (tashkeel)
        _claim("تم تعيين العامل في 2020", "factual_premise", "supporting"),
    ]})
    opinion = "نص الرأي: يستحق العامل إجازة ثلاثين يوماً، وقد تم تعيين العامل في 2020."
    claims, truncated = await extract_claims(llm, opinion, question="سؤال")
    assert [c.id for c in claims] == ["C1", "C2"]
    assert claims[0].cited_refs[0]["article"] == 70 and claims[1].type == "factual_premise"
    assert not truncated
    assert "سؤال" in llm.calls[0]["user"] and opinion in llm.calls[0]["user"]


@pytest.mark.asyncio
async def test_extract_claims_caps_at_max(monkeypatch) -> None:
    from app.verification import extract

    monkeypatch.setattr(extract, "MAX_CLAIMS", 2)
    llm = FakeLLM({"claims": [_claim(f"ادعاء رقم {i}") for i in range(5)]})
    claims, truncated = await extract_claims(llm, "ادعاء رقم " * 5, question=None)
    assert len(claims) == 2 and truncated


# ── validation ────────────────────────────────────────────────────────────

ART = "يستحق العامل إجازة سنوية مدفوعة الأجر مدتها ثلاثون يوماً ولا يجوز للعامل التنازل عنها."


def passage(pid: str, text: str = ART, **kw) -> Passage:
    base = dict(pid=pid, chunk_id=f"ch-{pid}", unit_id=f"u-{pid}", document_id="d", text=text,
                context_header="قانون: قانون العمل — المادة 70", title_ar="قانون العمل", number="6",
                year=2010, article_label="المادة 70", doc_type="law", authority=0.9, status="in_force",
                effective_date=date(2010, 2, 21), valid_from=date(2010, 2, 21), valid_to=None)
    return Passage(**(base | kw))


CLAIM = ClaimIn(id="C1", text_ar="للعامل إجازة سنوية ثلاثون يوماً", type="legal_conclusion",
                materiality="core", cited_refs=[])


def raw(verdict: str, *ev: tuple) -> dict:
    return {"claim_id": "C1", "verdict": verdict, "reasoning_ar": "تعليل",
            "evidence": [{"passage_id": p, "stance": s, "quote_ar": q} for p, s, q in ev]}


def test_verbatim_quote_is_kept_with_exact_source_span() -> None:
    r = validate(CLAIM, raw("supported", ("P1", "supports", "إجازة سنوية مدفوعة الاجر مدتها ثلاثون يوما")),
                 {"P1": passage("P1")}, date(2026, 1, 1))
    assert r.verdict == "supported"
    (e,) = r.evidence
    assert e.quote_ar == "إجازة سنوية مدفوعة الأجر مدتها ثلاثون يوماً"  # span from the source text
    assert e.chunk_id == "ch-P1"


def test_hallucinated_passage_and_invented_quote_are_dropped() -> None:
    r = validate(CLAIM, raw("supported", ("P9", "supports", "إجازة سنوية"),
                            ("P1", "supports", "يحق للعامل إجازة ستين يوماً في السنة كاملة")),
                 {"P1": passage("P1")}, date(2026, 1, 1))
    assert r.evidence == []
    assert {d["reason"] for d in r.dropped} == {"hallucinated_passage", "unverified_quote"}
    assert r.verdict == "insufficient"  # supported without remaining supports evidence


def test_repealed_source_cannot_support() -> None:
    r = validate(CLAIM, raw("supported", ("P1", "supports", "إجازة سنوية مدفوعة الأجر")),
                 {"P1": passage("P1", status="repealed")}, date(2026, 1, 1))
    assert r.evidence[0].stance == "context" and r.evidence[0].note
    assert r.verdict == "insufficient"


def test_contradicted_without_contradicting_evidence_becomes_insufficient() -> None:
    r = validate(CLAIM, raw("contradicted", ("P1", "context", "إجازة سنوية")), {"P1": passage("P1")},
                 date(2026, 1, 1))
    assert r.verdict == "insufficient"


def test_quotes_are_capped_at_40_words() -> None:
    long_text = " ".join(f"كلمة{i}" for i in range(60))
    r = validate(CLAIM, raw("supported", ("P1", "supports", long_text)), {"P1": passage("P1", text=long_text)},
                 date(2026, 1, 1))
    assert len(r.evidence[0].quote_ar.split()) == 40


# ── scoring ───────────────────────────────────────────────────────────────


def result(cid: str, verdict: str, typ: str = "legal_conclusion", mat: str = "core", ev=()):
    from app.verification.types import ClaimResult, Evidence

    c = ClaimIn(id=cid, text_ar=cid, type=typ, materiality=mat, cited_refs=[])
    return ClaimResult(claim=c, verdict=verdict, reasoning_ar="", dropped=[],
                       evidence=[Evidence(pid=p.pid, chunk_id=p.chunk_id, stance=s, quote_ar="q") for p, s in ev])


def test_weighted_score_excludes_factual_premises() -> None:
    p = {"P1": passage("P1")}
    rs = [result("C1", "supported", ev=[(p["P1"], "supports")]),              # core 2 × 1.0
          result("C2", "partially_supported", mat="supporting", ev=[(p["P1"], "supports")]),  # 1 × 0.5
          result("C3", "insufficient", mat="supporting"),                      # 1 × 0
          result("C4", "insufficient", typ="factual_premise")]                 # excluded
    out = score_review(rs, p, date(2026, 1, 1))
    assert out.score == round(100 * 2.5 / 4) and out.status == "needs_review"


def test_all_supported_is_accepted() -> None:
    p = {"P1": passage("P1")}
    out = score_review([result("C1", "supported", ev=[(p["P1"], "supports")])], p, date(2026, 1, 1))
    assert (out.score, out.status) == (100, "accepted")


def test_no_valid_evidence_means_no_information() -> None:
    out = score_review([result("C1", "insufficient"), result("C2", "insufficient", mat="supporting")], {},
                       date(2026, 1, 1))
    assert out.status == "no_information" and NO_INFO_MESSAGE.startswith("لا تتوفر معلومات")


def test_higher_authority_contradiction_blocks_lower_does_not() -> None:
    law = passage("P1", authority=0.9)
    circular = passage("P2", authority=0.5, doc_type="circular")
    const = passage("P3", authority=1.0, doc_type="constitution")
    p = {"P1": law, "P2": circular, "P3": const}
    # supported by a law, contradicted by a circular → resolved, not blocking
    r1 = result("C1", "supported", ev=[(law, "supports"), (circular, "contradicts")])
    out = score_review([r1], p, date(2026, 1, 1))
    assert out.status == "accepted"
    assert r1.evidence[1].blocking is False and r1.evidence[1].resolved_conflict is True
    # contradicted by the constitution → blocking even with a 100 score elsewhere
    r2 = result("C2", "supported", ev=[(law, "supports"), (const, "contradicts")])
    out = score_review([r2], p, date(2026, 1, 1))
    assert out.status == "needs_review" and r2.evidence[1].blocking is True


# ── hallucination guards (found by the live smoke test) ─────────────────────


def test_non_arabic_or_garbled_opinion_is_rejected_before_any_llm_call() -> None:
    from app.verification.parse import ensure_opinion_text

    with pytest.raises(ParseError) as e:
        ensure_opinion_text("??? ???????? ?? ?????? ???? 6 ?? ??????? ??? 6 ???? 2010. ??? ????")
    assert "عربي" in e.value.message_ar
    assert ensure_opinion_text("  يستحق العامل إجازة سنوية مدفوعة الأجر وفق القانون.  ") == (
        "يستحق العامل إجازة سنوية مدفوعة الأجر وفق القانون."
    )


@pytest.mark.asyncio
async def test_claims_not_grounded_in_the_opinion_are_dropped() -> None:
    opinion = "يستحق العامل إجازة سنوية مدفوعة الأجر مدتها ثلاثون يوماً وفق قانون العمل."
    llm = FakeLLM({"claims": [
        _claim("يستحق العامل إجازة سنوية مدفوعة الأجر ثلاثين يوماً"),
        _claim("يحق للمستأجر التأجير من الباطن بموافقة المؤجر الخطية"),  # invented
    ]})
    claims, _ = await extract_claims(llm, opinion, question=None)
    assert [c.text_ar for c in claims] == ["يستحق العامل إجازة سنوية مدفوعة الأجر ثلاثين يوماً"]


# ── final-review finding #1: quote validation must be verbatim and meaning-preserving ──

SRC_ALLOW = "يجوز لصاحب العمل إنهاء العقد بعد إخطار العامل كتابة قبل ثلاثة أشهر."
SRC_FORBID = "لا يجوز لصاحب العمل إنهاء عقد العامل غير محدد المدة إلا بعد إخطاره."


def _span(quote: str, source: str):
    from app.verification.validate import _source_span

    return _source_span(quote, source)


def test_added_negation_is_not_a_verbatim_quote() -> None:
    assert _span("لا يجوز لصاحب العمل إنهاء العقد", SRC_ALLOW) is None


def test_quote_with_invented_extra_clause_is_rejected() -> None:
    assert _span("يجوز لصاحب العمل إنهاء العقد بعد إخطار العامل مع دفع التعويض الكامل والمكافأة", SRC_ALLOW) is None


def test_one_word_quote_is_rejected() -> None:
    assert _span("العقد", SRC_ALLOW) is None


def test_dropped_negation_substring_is_rejected() -> None:
    assert _span("يجوز لصاحب العمل إنهاء عقد العامل غير محدد المدة", SRC_FORBID) is None


def test_genuine_quotes_still_pass_with_word_boundaries() -> None:
    assert _span("لا يجوز لصاحب العمل إنهاء عقد العامل", SRC_FORBID) == "لا يجوز لصاحب العمل إنهاء عقد العامل"
    # small OCR-ish typo tolerated, span snapped to whole source words
    assert _span("إنهاء العقد بعد اخطار العامل كتابه", SRC_ALLOW) == "إنهاء العقد بعد إخطار العامل كتابة"


# ── suspended articles are context only (user decision 2026-09-26) ──────────


def test_suspended_source_cannot_support() -> None:
    r = validate(CLAIM, raw("supported", ("P1", "supports", "إجازة سنوية مدفوعة الأجر مدتها ثلاثون يوماً")),
                 {"P1": passage("P1", status="suspended")}, date(2026, 1, 1))
    assert r.evidence[0].stance == "context" and r.verdict == "insufficient"


def test_suspended_contradiction_never_blocks() -> None:
    law = passage("P1", authority=0.7)
    const = passage("P2", authority=0.99, doc_type="constitution", status="suspended")
    r = result("C1", "supported", ev=[(law, "supports"), (const, "contradicts")])
    out = score_review([r], {"P1": law, "P2": const}, date(2026, 1, 1))
    assert r.evidence[1].blocking is False and out.status == "accepted"


# ── claims highlighted in the opinion (sentence span of each claim) ─────────

OPINION_TEXT = ("يرى المستشار أن العامل الذي أمضى في خدمة صاحب العمل سنة كاملة يستحق إجازة سنوية مدفوعة الأجر. "
                "كما يجوز الاتفاق على فترة تجربة للعامل لا تتجاوز ستة أشهر.\nوالله الموفق.")


def test_claim_is_located_on_its_sentence_in_the_opinion() -> None:
    from app.verification.validate import locate_claim

    s, e = locate_claim("يجوز الاتفاق على فترة تجربة لا تتجاوز ستة أشهر", OPINION_TEXT)
    assert OPINION_TEXT[s:e] == "كما يجوز الاتفاق على فترة تجربة للعامل لا تتجاوز ستة أشهر."
    s, e = locate_claim("العامل الذي أمضى سنة كاملة يستحق إجازة سنوية مدفوعة الأجر", OPINION_TEXT)
    assert OPINION_TEXT[s:e].startswith("يرى المستشار") and OPINION_TEXT[s:e].endswith("الأجر.")


def test_unrelated_claim_is_not_highlighted() -> None:
    from app.verification.validate import locate_claim

    assert locate_claim("يحق للمستأجر التأجير من الباطن بموافقة المؤجر الخطية", OPINION_TEXT) is None
